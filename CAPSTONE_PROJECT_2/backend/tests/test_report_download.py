"""
Test Suite: Report Generation & Binary Document Streaming
Validates ReportLab binary PDF generation and openpyxl multi-sheet Excel generation.
"""

import io
import unittest
from openpyxl import load_workbook
from backend.services.pdf_generator import generate_tourpulse_pdf
from backend.services.excel_generator import generate_tourpulse_excel

class TestReportDownloads(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        cls.test_report_data = {
            "city": "Chennai",
            "date_preset": "Last 30 Days",
            "kpis": {
                "total_visits": 1235,
                "active_attractions": 8,
                "top_attraction": "Marina Beach",
                "peak_visiting_hour": "17:00 (5:00 PM)",
                "avg_crowd_level_pct": 68.5,
                "daily_average_visitors": 41,
                "high_crowd_attractions_count": 2
            },
            "attractions": [
                {
                    "id": 1,
                    "name": "Marina Beach",
                    "city": "Chennai",
                    "category": "Coastal / Beach",
                    "current_visits": 340,
                    "capacity": 5000,
                    "current_crowd": "Moderate",
                    "popularity_index": 92
                },
                {
                    "id": 2,
                    "name": "Kapaleeshwarar Temple",
                    "city": "Chennai",
                    "category": "Heritage / Religious",
                    "current_visits": 220,
                    "capacity": 1500,
                    "current_crowd": "High",
                    "popularity_index": 88
                }
            ],
            "recommendations": [
                {
                    "name": "San Thome Basilica",
                    "city": "Chennai",
                    "category": "Heritage",
                    "current_crowd": "Low",
                    "optimal_time_slot": "10:00 AM – 01:00 PM",
                    "recommendation_reason": "Pairs with Kapaleeshwarar Temple for congestion relief."
                }
            ]
        }

    def test_pdf_generation_magic_bytes_and_length(self):
        pdf_bytes = generate_tourpulse_pdf(self.test_report_data)
        self.assertIsInstance(pdf_bytes, bytes)
        self.assertGreater(len(pdf_bytes), 1000, "PDF must not be empty or blank")
        self.assertTrue(pdf_bytes.startswith(b"%PDF"), "PDF binary must start with %PDF header")

    def test_excel_generation_multisheet(self):
        excel_bytes = generate_tourpulse_excel(self.test_report_data)
        self.assertIsInstance(excel_bytes, bytes)
        self.assertGreater(len(excel_bytes), 1000, "Excel file must not be empty")
        self.assertTrue(excel_bytes.startswith(b"PK"), "XLSX must be a valid zip archive")

        # Load workbook and verify sheets
        wb = load_workbook(io.BytesIO(excel_bytes))
        sheet_names = wb.sheetnames
        self.assertIn("Summary", sheet_names, "Workbook must have 'Summary' sheet")
        self.assertIn("Attractions", sheet_names, "Workbook must have 'Attractions' sheet")
        self.assertIn("Recommendations", sheet_names, "Workbook must have 'Recommendations' sheet")

        # Verify Summary content
        ws_sum = wb["Summary"]
        self.assertEqual(ws_sum["B5"].value, "Chennai")
        self.assertEqual(ws_sum["B12"].value, 1235)

if __name__ == "__main__":
    unittest.main()

